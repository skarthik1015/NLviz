"""Uploads served through Jammi: conversion, managed storage, and the upload/delete routes."""
from __future__ import annotations

from pathlib import Path

import pyarrow as pa
import pyarrow.parquet as pq
import pytest

from app.dependencies import get_jammi_provisioner, get_jammi_upload_store
from app.services.jammi_sources import JammiSourceProvisioner, source_id_for, workspace_tag
from app.services.jammi_uploads import (
    LocalUploadStore,
    S3UploadStore,
    UploadError,
    convert_to_parquet,
    table_name_for,
)

from test_connections_route import _TEST_USER, _build_connection_app  # tests/ is on sys.path under pytest
from test_jammi_connection_routes import FakeJammi

CSV = b"Order ID,Customer Name,order id,2024 Total,\n1,Ana,10,1.5,x\n2,Bruno,11,2.5,y\n"
CONNECTION_ID = "123e4567-e89b-42d3-a456-556642440000"


# ── Conversion ───────────────────────────────────────────────────────


def test_csv_is_converted_with_queryable_column_names(tmp_path: Path):
    source = tmp_path / "in.csv"
    source.write_bytes(CSV)

    result = convert_to_parquet(source, ".csv", tmp_path / "out" / "sales.parquet")

    table = pq.read_table(result.path)
    assert table.column_names == ["order_id", "customer_name", "order_id_2", "c_2024_total", "column_5"]
    assert result.columns == table.column_names and result.num_rows == 2
    assert result.renamed_columns == {
        "Order ID": "order_id", "Customer Name": "customer_name", "order id": "order_id_2",
        "2024 Total": "c_2024_total", "": "column_5",
    }
    assert table.column("customer_name").to_pylist() == ["Ana", "Bruno"]


def test_parquet_is_rewritten_and_lowercase_names_are_kept(tmp_path: Path):
    source = tmp_path / "in.parquet"
    pq.write_table(pa.table({"id": [1, 2], "amount": [3.0, 4.0]}), source)

    result = convert_to_parquet(source, ".parquet", tmp_path / "t.parquet")

    assert result.columns == ["id", "amount"] and result.renamed_columns == {}
    assert pq.read_table(result.path).column("amount").to_pylist() == [3.0, 4.0]


@pytest.mark.parametrize("suffix,payload", [
    (".parquet", b"not parquet"),
    (".csv", b"a,b\n1,2,3\n"),  # ragged row
    (".csv", "a,b\n1,2\n".encode("utf-16")),
    (".csv", b"a,b\n"),  # header only: nothing to query
])
def test_unreadable_files_raise_a_user_facing_error(tmp_path: Path, suffix, payload):
    source = tmp_path / f"in{suffix}"
    source.write_bytes(payload)
    with pytest.raises(UploadError):
        convert_to_parquet(source, suffix, tmp_path / "t.parquet")


def test_excel_first_sheet_with_data_is_read_as_typed_columns(tmp_path: Path):
    import datetime

    import openpyxl

    book = openpyxl.Workbook()
    notes = book.active
    notes.title = "Notes"  # empty: skipped
    data = book.create_sheet("Sales")
    data.append([])  # a blank row above the header
    data.append(["Region", "Amount", "Code", "Sold On", None])
    data.append(["North", 10.5, 7, datetime.datetime(2024, 1, 5), None])
    data.append(["South", 3, "N/A", datetime.datetime(2024, 2, 1), None])
    data.append([None, None, None, None, None])
    book.create_sheet("Pivot").append(["x"])
    source = tmp_path / "in.xlsx"
    book.save(source)

    out = convert_to_parquet(source, ".xlsx", tmp_path / "t.parquet")
    table = pq.read_table(out.path)
    assert out.columns == ["region", "amount", "code", "sold_on"] and out.num_rows == 2
    assert pa.types.is_floating(table.schema.field("amount").type)
    assert pa.types.is_timestamp(table.schema.field("sold_on").type)
    assert table.column("code").to_pylist() == ["7", "N/A"]  # numbers and text: text
    assert out.notes and "'Sales'" in out.notes[0] and "Notes, Pivot" in out.notes[0]


def test_a_file_that_is_not_a_workbook_is_refused(tmp_path: Path):
    source = tmp_path / "in.xlsx"
    source.write_bytes(b"region,amount\nNorth,1\n")
    with pytest.raises(UploadError, match="Excel"):
        convert_to_parquet(source, ".xlsx", tmp_path / "t.parquet")


@pytest.mark.parametrize("payload", [
    "city,n\nSão Paulo,1\n".encode("cp1252"),  # Windows-1252 data: pyarrow would type the column binary
    "Município,n\nSão Paulo,1\n".encode("cp1252"),  # ...and header: fails only when the name is read
])
def test_windows_1252_csv_is_read_as_text(tmp_path: Path, payload):
    source = tmp_path / "in.csv"
    source.write_bytes(payload)

    table = pq.read_table(convert_to_parquet(source, ".csv", tmp_path / "t.parquet").path)

    assert pa.types.is_string(table.schema.field(0).type)
    assert table.column(0).to_pylist() == ["São Paulo"]
    assert table.column_names[0] in ("city", "munic_pio")


@pytest.mark.parametrize("filename,table", [
    ("Sales 2024.csv", "sales_2024"),
    ("2024-sales.parquet", "t_2024_sales"),
    ("C:\\Users\\me\\Q1 Report.CSV", "q1_report"),
    ("%%%.csv", "uploaded_table"),
])
def test_table_names_are_lowercase_identifiers(filename, table):
    assert table_name_for(filename) == table


# ── Managed storage ──────────────────────────────────────────────────


def test_local_store_writes_under_a_generated_workspace_path(tmp_path: Path):
    store = LocalUploadStore(tmp_path / "uploads", "/uploads")
    local = tmp_path / "x.parquet"
    local.write_bytes(b"PAR1")

    url = store.put(local, workspace_id="user-1", connection_id=CONNECTION_ID, table="sales")

    expected = f"w{workspace_tag('user-1')}/123e4567e89b42d3a456556642440000/sales.parquet"
    assert url == f"/uploads/{expected}"
    assert (tmp_path / "uploads" / expected).read_bytes() == b"PAR1"

    store.delete(url)
    assert not (tmp_path / "uploads" / expected).exists()
    assert list((tmp_path / "uploads").iterdir()) == []  # empty connection/workspace dirs removed


@pytest.mark.parametrize("url", ["/elsewhere/x.parquet", "/uploads/../secrets/x.parquet", "/uploads/../../etc/passwd"])
def test_local_store_refuses_to_delete_outside_its_root(tmp_path: Path, url):
    store = LocalUploadStore(tmp_path / "uploads", "/uploads")
    with pytest.raises(ValueError):
        store.delete(url)


def test_local_store_rejects_unsafe_table_names(tmp_path: Path):
    store = LocalUploadStore(tmp_path / "uploads", "/uploads")
    with pytest.raises(ValueError):
        store.put(tmp_path / "x", workspace_id="u", connection_id=CONNECTION_ID, table="../escape")


def test_s3_store_uses_its_prefix_and_only_deletes_inside_it(tmp_path: Path):
    calls = []

    class FakeS3:
        bucket = "nlviz-uploads"

        def upload_fileobj(self, key, fobj):
            calls.append(("put", key, fobj.read()))
            return f"s3://{self.bucket}/{key}"

        def key_from_url(self, url):
            return url.split(f"s3://{self.bucket}/", 1)[1]

        def delete(self, key):
            calls.append(("delete", key))

    local = tmp_path / "x.parquet"
    local.write_bytes(b"PAR1")
    store = S3UploadStore(FakeS3())
    url = store.put(local, workspace_id="user-1", connection_id=CONNECTION_ID, table="sales")

    key = f"jammi-uploads/w{workspace_tag('user-1')}/123e4567e89b42d3a456556642440000/sales.parquet"
    assert url == f"s3://nlviz-uploads/{key}" and calls == [("put", key, b"PAR1")]
    store.delete(url)
    assert calls[-1] == ("delete", key)
    with pytest.raises(ValueError):
        store.delete("s3://nlviz-uploads/schemas/other.yaml")


# ── Routes ───────────────────────────────────────────────────────────


def _app(tmp_path: Path, jammi: FakeJammi):
    client, store, secrets = _build_connection_app(tmp_path)
    uploads = LocalUploadStore(tmp_path / "uploads", "/uploads")
    client.app.dependency_overrides[get_jammi_provisioner] = lambda: JammiSourceProvisioner(jammi)
    client.app.dependency_overrides[get_jammi_upload_store] = lambda: uploads
    return client, store, secrets


def _upload(client, name="Sales 2024.csv", payload=CSV):
    return client.post("/connections/upload", files={"file": (name, payload, "text/csv")},
                       data={"display_name": "Sales"})


def test_upload_is_registered_as_one_jammi_source(tmp_path: Path):
    jammi = FakeJammi()
    client, store, secrets = _app(tmp_path, jammi)

    response = _upload(client)

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["renamed_columns"]["Order ID"] == "order_id"
    profile = store.get_connection(body["connection_id"])
    sid = source_id_for(_TEST_USER.user_id, body["connection_id"], "sales_2024")
    assert profile.engine == "jammi" and profile.connector_type == "upload"
    assert profile.workspace_id == _TEST_USER.user_id
    assert profile.jammi_tables == {"sales_2024": f"{sid}.public.sales_2024"}
    url = jammi.sources[sid]
    assert url.startswith("/uploads/") and url.endswith("/sales_2024.parquet")
    assert secrets.get_by_connection_id(body["connection_id"]) == {"upload_urls": [url]}
    stored = tmp_path / "uploads" / url.removeprefix("/uploads/")
    assert pq.read_table(stored).num_rows == 2


def test_unreadable_upload_is_refused_and_leaves_nothing(tmp_path: Path):
    jammi = FakeJammi()
    client, store, secrets = _app(tmp_path, jammi)

    response = _upload(client, name="broken.parquet", payload=b"not parquet")

    assert response.status_code == 400
    assert response.json()["detail"].startswith("UPLOAD_INVALID")
    assert jammi.sources == {} and store.list_connections(owner_id=_TEST_USER.user_id) == []
    assert not (tmp_path / "uploads").exists() or not any((tmp_path / "uploads").rglob("*.parquet"))


def test_registration_failure_deletes_the_stored_file(tmp_path: Path):
    jammi = FakeJammi(fail_add_with="InvalidArgument: no files matched extension '.parquet' under '/uploads/...'")
    client, store, secrets = _app(tmp_path, jammi)

    response = _upload(client)

    assert response.status_code == 400
    detail = response.json()["detail"]
    assert detail.startswith("JAMMI_REGISTRATION_FAILED") and "/uploads/" not in detail
    assert not any((tmp_path / "uploads").rglob("*.parquet"))
    assert store.list_connections(owner_id=_TEST_USER.user_id) == []
    assert not any((tmp_path / ".secrets").glob("*"))


def test_store_failure_removes_source_file_and_params(tmp_path: Path, monkeypatch):
    jammi = FakeJammi()
    client, store, secrets = _app(tmp_path, jammi)

    def broken_create(profile):
        raise RuntimeError("database down")

    monkeypatch.setattr(store, "create_connection", broken_create)
    with pytest.raises(RuntimeError, match="database down"):
        _upload(client)

    assert jammi.sources == {} and len(jammi.removed) == 1
    assert not any((tmp_path / "uploads").rglob("*.parquet"))
    assert not any((tmp_path / ".secrets").glob("*"))


def test_delete_removes_the_source_and_the_file(tmp_path: Path):
    jammi = FakeJammi()
    client, store, secrets = _app(tmp_path, jammi)
    connection_id = _upload(client).json()["connection_id"]
    assert any((tmp_path / "uploads").rglob("*.parquet"))

    class StubConnectionService:
        def evict_runtime(self, connection_id: str) -> None:
            pass

    from app.dependencies import get_connection_service

    client.app.dependency_overrides[get_connection_service] = lambda: StubConnectionService()
    response = client.delete(f"/connections/{connection_id}")

    assert response.status_code == 200
    assert jammi.sources == {} and len(jammi.removed) == 1
    assert not any((tmp_path / "uploads").rglob("*.parquet"))


def test_without_upload_storage_uploads_are_refused(tmp_path: Path):
    jammi = FakeJammi()
    client, store, _ = _build_connection_app(tmp_path)
    client.app.dependency_overrides[get_jammi_provisioner] = lambda: JammiSourceProvisioner(jammi)

    response = _upload(client)

    assert response.status_code == 400 and response.json()["detail"].startswith("JAMMI_NOT_CONFIGURED")
    assert jammi.sources == {} and store.list_connections(owner_id=_TEST_USER.user_id) == []
